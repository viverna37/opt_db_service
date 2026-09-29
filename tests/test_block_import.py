"""Импорт «блочного» прайса (формат Amigo Opt): парсер + синхронизация каталога"""
import io
from argparse import Namespace

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill
from PIL import Image

import app.database.models as m
from app.importer.block_price import parse_workbook
from tests.helpers import headers, run, seed_member, seed_tenant

ORANGE, BRAND, LIQ, TIER = "FF9900", "F9CB9C", "E69138", "F9CB9C"


def _fill(color):
    return PatternFill("solid", fgColor=color)


def _png():
    buf = io.BytesIO()
    Image.new("RGB", (40, 40), (10, 200, 10)).save(buf, format="PNG")
    buf.seek(0)
    return buf


def build_price(ribbon_price=510, with_second=True) -> bytes:
    wb = Workbook()
    about = wb.active
    about.title = "О НАС!"
    about["A1"] = "КОМПАНИЯ"

    # Вертикальные уровни + бренд-заливка + вкусы списком в одной ячейке
    ws = wb.create_sheet("ЭЛЕКТРОНКИ")
    ws["C1"] = "КОМПАНИЯ AMIGO OPT"
    ws["A15"] = "ПРОДУКЦИЯ  ICON"; ws["A15"].fill = _fill("F6B26B")
    ws["A17"] = "ICON 40000 "; ws["A17"].fill = _fill(ORANGE)
    ws["F17"] = "ГРАДАЦИЯ ЦЕН"; ws["H17"] = "КРУПНЫЙ ОБЪЁМ ОБСУЖДАЕТСЯ"
    ws["A19"] = "1. Hawaii blueberry ice\n2. Ice mint\n3.⁠ ⁠Mexican mango"
    rows = [("ОТ 3.000", ribbon_price), ("ОТ 10.000", 500), ("ОТ 300.000", "индивидуальное предложение"), ("РОЗНИЦА ОТ 1000", 1100)]
    for i, (label, price) in enumerate(rows):
        ws.cell(row=19 + i * 2, column=6, value=label)
        ws.cell(row=19 + i * 2, column=7, value=price)
    ws["H19"] = "НОВИНКА"
    img = XLImage(_png()); img.anchor = "C17"; ws.add_image(img)
    if with_second:
        # зачёркнутый товар «строго микс»
        ws["A40"] = "SOAK 5000"; ws["A40"].fill = _fill(ORANGE); ws["A40"].font = Font(strike=True)
        ws["I40"] = "СТРОГО МИКС (можете отметить предпочтение)"
        ws["F42"] = "ОТ 3.000"; ws["G42"] = 220
        ws["F44"] = "РОЗНИЦА ОТ 1000"; ws["G44"] = 550

    # Горизонтальные уровни, бренд — строка без заливки, вкусы колонками, зачёркнутый вкус
    lq = wb.create_sheet("ЖИДКОСТИ")
    lq["A15"] = "ЖНЕЦ"
    lq["A17"] = "ЖНЕЦ 30ML|50MG "; lq["A17"].fill = _fill(LIQ)
    for col, label in zip("BCD", ("ОТ 3.000", "ОТ 30.000Р", "ОТ 50.000Р")):
        lq[f"{col}17"] = label; lq[f"{col}17"].fill = _fill(TIER)
    for col, price in zip("BCD", (260, 255, 245)):
        lq[f"{col}19"] = price
    lq["A22"] = "Мятная жвачка айс"; lq["C22"] = "Кола айс"
    lq["A23"] = "Манго айс"; lq["A23"].font = Font(strike=True)
    lq["A24"] = "Арбуз"  # вплотную к списку — вкус, а не бренд следующего товара
    lq["A26"] = "OGGO 30ml 50mg"; lq["A26"].fill = _fill(LIQ)
    lq["B26"] = "ОТ 3.000"; lq["B28"] = 260
    lq["A30"] = "15 вкусов"

    hidden = wb.create_sheet("СЛИВ")
    hidden.sheet_state = "hidden"
    hidden["A17"] = "СКРЫТЫЙ"; hidden["A17"].fill = _fill(ORANGE); hidden["F19"] = "ОТ 3.000"; hidden["G19"] = 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parser_reads_blocks():
    res = parse_workbook(build_price())
    assert [s.title for s in res.sheets] == ["ЭЛЕКТРОНКИ", "ЖИДКОСТИ"]  # «О НАС» и скрытые пропущены
    icon, soak = res.sheets[0].products
    assert icon.name == "ICON 40000" and icon.brand == "ICON" and icon.image
    assert icon.prices == {3000: 51000, 10000: 50000, 1000: 110000}  # «индивидуальное предложение» пропущено
    assert [v.name for v in icon.variants] == ["Hawaii blueberry ice", "Ice mint", "Mexican mango"]
    assert soak.out and soak.mix and soak.variants == [] and soak.brand is None

    reaper, oggo = res.sheets[1].products
    assert reaper.brand == "ЖНЕЦ" and reaper.prices == {3000: 26000, 30000: 25500, 50000: 24500}
    assert [(v.name, v.out) for v in reaper.variants] == [
        ("Мятная жвачка айс", False), ("Кола айс", False), ("Манго айс", True), ("Арбуз", False),
    ]
    assert oggo.brand is None  # «Арбуз» — хвост списка вкусов, а не бренд
    assert oggo.description == "15 вкусов" and oggo.variants == []


def test_import_creates_catalog_and_resyncs(client, session_factory, tmp_path):
    from app.cli import _with_repo, import_price

    tenant_id = seed_tenant(session_factory, "amigo", age_gate=False)
    seed_member(session_factory, tenant_id, 1, role=m.Role.owner)
    owner = headers("amigo", 1)
    # тестовый мусор, который --wipe должен убрать
    client.post("/v1/admin/price-tiers", json={"label": "от 5 шт", "min_qty": 5}, headers=owner)
    client.post("/v1/admin/products", json={"name": "Скала"}, headers=owner)

    path = tmp_path / "price.xlsx"
    path.write_bytes(build_price())
    run(_with_repo(import_price, Namespace(slug="amigo", file=str(path), wipe=True, no_photos=False, dry_run=False)))

    me = client.get("/v1/me", headers=owner).json()
    assert me["tenant"]["price_basis"] == "amount" and me["tenant"]["min_order_amount"] == 100000
    tiers = client.get("/v1/admin/price-tiers", headers=owner).json()
    assert [(t["label"], t["min_amount"]) for t in tiers] == [
        ("Розница от 1 000 ₽", 100000), ("от 3 000 ₽", 300000), ("от 10 000 ₽", 1000000),
        ("от 30 000 ₽", 3000000), ("от 50 000 ₽", 5000000),
    ]
    assert [c["name"] for c in client.get("/v1/catalog/categories", headers=owner).json()] == ["Электронки", "Жидкости"]
    products = client.get("/v1/admin/products", headers=owner).json()["items"]
    assert sorted(p["name"] for p in products) == ["ICON 40000", "OGGO 30ml 50mg", "SOAK 5000", "ЖНЕЦ 30ML|50MG"]

    icon = next(p for p in products if p["name"] == "ICON 40000")
    card = client.get(f"/v1/catalog/products/{icon['id']}", headers=owner).json()
    assert card["brand"] == "ICON" and len(card["photos"]) == 1 and len(card["variants"]) == 3
    assert [t["amount"] for t in card["tiers"]] == [110000, 51000, 50000, 50000, 50000]

    reaper = next(p for p in products if p["name"].startswith("ЖНЕЦ"))
    card = client.get(f"/v1/catalog/products/{reaper['id']}", headers=owner).json()
    assert {v["name"]: v["stock_status"] for v in card["variants"]}["Манго айс"] == "out"
    soak = next(p for p in products if p["name"] == "SOAK 5000")
    assert soak["stock_status"] == "out" and soak["variants_count"] == 0

    # Клиент кладёт вариант в корзину — повторный импорт не должен его ломать
    client.put(f"/v1/cart/items/{card['variants'][0]['id']}", json={"qty": 2}, headers=headers("amigo", 10))

    # Следующий прайс: цена ICON изменилась, SOAK пропал; фото не дублируется
    path.write_bytes(build_price(ribbon_price=490, with_second=False))
    run(_with_repo(import_price, Namespace(slug="amigo", file=str(path), wipe=False, no_photos=False, dry_run=False)))
    products = {p["name"]: p for p in client.get("/v1/admin/products", headers=owner).json()["items"]}
    assert products["SOAK 5000"]["is_visible"] is False
    assert products["ICON 40000"]["id"] == icon["id"]
    card = client.get(f"/v1/catalog/products/{icon['id']}", headers=owner).json()
    assert card["tiers"][1]["amount"] == 49000 and len(card["photos"]) == 1
    assert client.get("/v1/cart", headers=headers("amigo", 10)).json()["total_qty"] == 2
