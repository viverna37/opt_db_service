"""Парсер витринного прайса «товар = блок строк» (формат «Сафари вейп»)"""
from io import BytesIO

import openpyxl
from openpyxl.styles import PatternFill

from app.importer.detect import parse_price_file
from app.importer.merged_price import parse_merged

RED = PatternFill(fill_type="solid", fgColor="FFFF0000")


def _workbook() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "СнюсПластинки"
    ws["A2"] = "Если строка выделена красным - нет в наличии"
    ws.append([])
    ws["A4"], ws["B4"], ws["C4"], ws["D4"] = "Фото товара", "Название", "Вкусы", "Цена"
    ws["A5"] = "Vaporesso (к каждому устройству жижа в подарок🎁)"
    ws.merge_cells("A5:D5")
    # карточка 1: фото A6:A10, три вкуса, один нет в наличии, скидка от 2 шт
    ws.merge_cells("A6:A10")
    ws["B6"], ws["C6"], ws["D6"] = "XROS 5", "Синий", 2500
    ws["C7"] = "Розовый"
    ws["C8"] = "Белый"
    ws["C8"].fill = RED
    ws["D9"] = "(ПРИ ПОКУПКЕ ОТ 2 шт"
    ws["D10"] = "Цена будет 2300₽"
    # карточка 2: название в две строки, цена строкой «350р», без вкусов, вся красная
    ws.merge_cells("A11:A14")
    ws["B11"], ws["D11"] = "NEW\n\nFEELIN 2", "350р"
    ws["B11"].fill = RED
    # карточка 3: без объединения в A — начало по объединению названия
    ws.merge_cells("B15:B16")
    ws["B15"], ws["C15"], ws["D15"] = "SANTI", "Чёрный", 1790
    return _save(wb)


def _save(wb) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_blocks_brand_stock_and_bulk_price():
    result = parse_merged(_workbook())
    assert [s.title for s in result.sheets] == ["Снюс, пластинки"]
    xros, feelin, santi = result.products
    assert santi.name == "SANTI" and [v.name for v in santi.variants] == ["Чёрный"]
    assert xros.brand == "Vaporesso"
    assert xros.prices == {1: 250000, 2: 230000}
    assert [(v.name, v.out) for v in xros.variants] == [("Синий", False), ("Розовый", False), ("Белый", True)]
    assert not xros.out
    assert "От 2 шт — 2300 ₽" in xros.description and "жижа в подарок" in xros.description
    assert feelin.name == "FEELIN 2" and feelin.prices == {1: 35000}
    assert feelin.out and "Новинка" in feelin.description
    assert result.tenant_defaults == {"price_basis": "qty"}


def test_detected_as_merged():
    kind, result = parse_price_file(_workbook(), "safari.xlsx")
    assert kind == "merged" and len(result.products) == 3
