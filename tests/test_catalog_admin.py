"""Каталог (дерево, наследование атрибутов, фильтры, поиск) и админка (заявки, живые корзины, фото, справочники)"""
import io
from argparse import Namespace

from openpyxl import load_workbook
from PIL import Image

import app.database.models as m
from tests.helpers import (
    create_product,
    headers,
    notifications,
    run,
    seed_member,
    setup_shop,
    visible_variant_ids,
)


def _catalog(client, session_factory):
    ctx = setup_shop(client, session_factory, age_gate=False)
    owner = ctx["owner"]
    volume = client.post("/v1/admin/attributes", json={
        "key": "volume", "label": "Объём", "type": "number", "unit": "мл", "filterable": True, "show_in_list": True,
    }, headers=owner).json()
    nicotine = client.post("/v1/admin/attributes", json={
        "key": "nicotine", "label": "Никотин", "type": "select", "unit": "мг", "options": ["20", "50"],
        "filterable": True,
    }, headers=owner).json()
    color = client.post("/v1/admin/attributes", json={
        "key": "color", "label": "Цвет", "type": "color", "scope": "variant", "filterable": True,
    }, headers=owner).json()
    liquids = client.post("/v1/admin/categories", json={"name": "Жидкости"}, headers=owner).json()
    salt = client.post("/v1/admin/categories", json={"name": "Солевые", "parent_id": liquids["id"]},
                       headers=owner).json()
    pods = client.post("/v1/admin/categories", json={"name": "Поды", "sort_order": 1}, headers=owner).json()
    client.put(f"/v1/admin/categories/{liquids['id']}/attributes",
               json={"attribute_ids": [volume["id"], nicotine["id"]]}, headers=owner)
    client.put(f"/v1/admin/categories/{pods['id']}/attributes", json={"attribute_ids": [color["id"]]}, headers=owner)

    p30 = create_product(client, ctx, name="Жнец 30", category_id=salt["id"],
                         attributes={"volume": "30", "nicotine": "50"}, prices=(26000, 25000, 24000))
    p60 = create_product(client, ctx, name="Ёлка 60", category_id=salt["id"],
                         attributes={"volume": 60, "nicotine": "20"}, prices=(30000, 29000, 28000))
    pod = create_product(client, ctx, name="Charon Baby", category_id=pods["id"], variants=(),
                         prices=(75000, 73000, 70000))
    client.post(f"/v1/admin/products/{pod['id']}/variants", json={"name": "Black", "attributes": {"color": "#000"}},
                headers=owner)
    client.post(f"/v1/admin/products/{pod['id']}/variants", json={"name": "Blue", "attributes": {"color": "#00F"}},
                headers=owner)
    return ctx, {"liquids": liquids, "salt": salt, "pods": pods}, {"p30": p30, "p60": p60, "pod": pod}


def test_category_tree_inheritance_and_filters(client, session_factory):
    ctx, cats, products = _catalog(client, session_factory)
    h = headers("shop", 10)
    tree = client.get("/v1/catalog/categories", headers=h).json()
    assert [(n["name"], n["product_count"]) for n in tree] == [("Жидкости", 2), ("Поды", 1)]
    assert tree[0]["children"][0]["name"] == "Солевые"

    detail = client.get(f"/v1/catalog/categories/{cats['salt']['id']}", headers=h).json()
    assert [b["name"] for b in detail["breadcrumbs"]] == ["Жидкости", "Солевые"]
    salt_url = f"/v1/admin/categories/{cats['salt']['id']}/attributes"
    assert client.get(salt_url, headers=ctx["owner"]).json() == []
    assert len(client.get(f"{salt_url}?inherited=true", headers=ctx["owner"]).json()) == 2
    # атрибуты унаследованы от «Жидкости»; значения фильтров — из товаров / options
    filters = {f["attribute"]["key"]: f["values"] for f in detail["filters"]}
    assert filters == {"volume": [30, 60], "nicotine": ["20", "50"]}

    # number хранится числом — фильтр работает и для "30", пришедшего строкой в админке
    r = client.get(f"/v1/catalog/products?category_id={cats['liquids']['id']}&attr.volume=30", headers=h).json()
    assert [p["name"] for p in r["items"]] == ["Жнец 30"]
    assert r["items"][0]["meta"][0] == {"key": "volume", "label": "Объём", "value": 30, "unit": "мл",
                                        "type": "number"}
    r = client.get("/v1/catalog/products?attr.nicotine=20,50&sort=price_desc", headers=h).json()
    assert [p["name"] for p in r["items"]] == ["Ёлка 60", "Жнец 30"]
    # фильтр по атрибуту варианта
    r = client.get("/v1/catalog/products?attr.color=%2300F", headers=h).json()
    assert [p["name"] for p in r["items"]] == ["Charon Baby"]
    # неизвестный ключ игнорируется
    assert client.get("/v1/catalog/products?attr.nope=1", headers=h).json()["total"] == 3


def test_search_and_in_stock(client, session_factory):
    ctx, cats, products = _catalog(client, session_factory)
    h = headers("shop", 10)
    # ё/е и регистр не важны; ищется и по названию варианта
    assert [p["name"] for p in client.get("/v1/catalog/products?q=елка", headers=h).json()["items"]] == ["Ёлка 60"]
    assert [p["name"] for p in client.get("/v1/catalog/products?q=blue", headers=h).json()["items"]] == ["Charon Baby"]

    r = client.post(f"/v1/admin/products/{products['p30']['id']}/stock", json={"stock_status": "out"},
                    headers=ctx["owner"])
    assert all(v["stock_status"] == "out" for v in r.json()["variants"] if v["is_visible"])
    names = [p["name"] for p in client.get("/v1/catalog/products?in_stock=true", headers=h).json()["items"]]
    assert "Жнец 30" not in names and len(names) == 2
    assert client.get("/v1/admin/summary", headers=ctx["owner"]).json()["out_of_stock_products"] == 1


def test_attribute_validation(client, session_factory):
    ctx, cats, products = _catalog(client, session_factory)
    owner = ctx["owner"]
    pid = products["p30"]["id"]
    assert client.patch(f"/v1/admin/products/{pid}", json={"attributes": {"volume": "abc"}},
                        headers=owner).status_code == 422
    assert client.patch(f"/v1/admin/products/{pid}", json={"attributes": {"nicotine": "99"}},
                        headers=owner).status_code == 422
    assert client.patch(f"/v1/admin/products/{pid}", json={"attributes": {"unknown": 1}},
                        headers=owner).status_code == 422
    # атрибут варианта нельзя задать товару
    assert client.patch(f"/v1/admin/products/{pid}", json={"attributes": {"color": "#fff"}},
                        headers=owner).status_code == 422
    assert client.post("/v1/admin/attributes", json={"key": "volume", "label": "x", "type": "text"},
                       headers=owner).status_code == 409


def test_hidden_category_and_product(client, session_factory):
    ctx, cats, products = _catalog(client, session_factory)
    h, owner = headers("shop", 10), ctx["owner"]
    client.patch(f"/v1/admin/categories/{cats['pods']['id']}", json={"is_visible": False}, headers=owner)
    assert [n["name"] for n in client.get("/v1/catalog/categories", headers=h).json()] == ["Жидкости"]
    assert client.get(f"/v1/catalog/categories/{cats['pods']['id']}", headers=h).status_code == 404

    client.delete(f"/v1/admin/products/{products['p60']['id']}", headers=owner)
    assert client.get(f"/v1/catalog/products/{products['p60']['id']}", headers=h).status_code == 404
    # категорию с товарами/подкатегориями не удалить; цикл в дереве не создать
    assert client.delete(f"/v1/admin/categories/{cats['liquids']['id']}", headers=owner).status_code == 409
    r = client.patch(f"/v1/admin/categories/{cats['liquids']['id']}", json={"parent_id": cats["salt"]["id"]},
                     headers=owner)
    assert r.status_code == 422


def test_stock_qty_and_threshold(client, session_factory):
    ctx = setup_shop(client, session_factory, age_gate=False)
    product = create_product(client, ctx, variants=("A",))
    variant = visible_variant_ids(product)[0]
    url = f"/v1/admin/products/{product['id']}/variants/{variant}"
    v = client.patch(url, json={"stock_qty": 4}, headers=ctx["owner"]).json()["variants"]
    assert [(x["stock_qty"], x["stock_status"]) for x in v if x["id"] == variant] == [(4, "low")]
    client.patch("/v1/admin/settings", json={"low_stock_threshold": 2}, headers=ctx["owner"])
    v = client.get(f"/v1/admin/products/{product['id']}", headers=ctx["owner"]).json()["variants"]
    assert [x["stock_status"] for x in v if x["id"] == variant] == ["in_stock"]
    v = client.patch(url, json={"stock_status": "out"}, headers=ctx["owner"]).json()["variants"]
    assert [(x["stock_qty"], x["stock_status"]) for x in v if x["id"] == variant] == [(None, "out")]


def test_admin_order_flow(client, session_factory):
    ctx = setup_shop(client, session_factory, age_gate=False)
    product = create_product(client, ctx, sku="ZH-30")
    mango = visible_variant_ids(product)[0]
    h = headers("shop", 10)
    client.put(f"/v1/cart/items/{mango}", json={"qty": 3}, headers=h)
    order = client.post("/v1/cart/submit", json={"comment": "срочно"}, headers=h).json()
    owner = ctx["owner"]

    summary = client.get("/v1/admin/summary", headers=owner).json()
    assert summary["new_orders"] == 1 and summary["live_carts"] == 0

    detail = client.get(f"/v1/admin/orders/{order['id']}", headers=owner).json()
    assert detail["client"]["contact_url"] == "https://t.me/user10" or detail["client"]["contact_url"].startswith("tg://")
    r = client.patch(f"/v1/admin/orders/{order['id']}/status", json={"status": "in_progress"}, headers=owner).json()
    assert r["status"] == "in_progress" and [x["to_status"] for x in r["history"]] == ["new", "in_progress"]
    # повторное «Взять в работу» из бота не плодит историю
    r = client.patch(f"/v1/admin/orders/{order['id']}/status", json={"status": "in_progress"}, headers=owner).json()
    assert len(r["history"]) == 2
    r = client.patch(f"/v1/admin/orders/{order['id']}/note", json={"manager_note": "позвонить"}, headers=owner)
    assert r.json()["manager_note"] == "позвонить"
    assert "manager_note" not in client.get(f"/v1/orders/{order['id']}", headers=h).json()

    assert client.get("/v1/admin/orders?status=new", headers=owner).json()["total"] == 0
    assert client.get("/v1/admin/orders?status=in_progress", headers=owner).json()["total"] == 1

    text = client.get(f"/v1/admin/orders/{order['id']}/text", headers=owner).text
    assert "Заявка №1" in text and "Манго [ZH-30]: 3 шт × 260 ₽ = 780 ₽" in text and "Итого: 780 ₽" in text

    xlsx = client.get(f"/v1/admin/orders/{order['id']}/export.xlsx", headers=owner)
    sheet = load_workbook(io.BytesIO(xlsx.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][0] == "№ заявки" and rows[1][4] == "Жнец 30мл" and rows[1][7] == 3 and rows[1][10] == 780
    assert client.get("/v1/admin/orders/export.xlsx?status=in_progress", headers=owner).status_code == 200

    audit = client.get("/v1/admin/audit", headers=owner).json()
    assert any(a["entity"] == "order" and a["action"] == "status" for a in audit)


def test_live_carts_and_remind_once_a_day(client, session_factory):
    ctx = setup_shop(client, session_factory, age_gate=False)
    product = create_product(client, ctx)
    mango, cola = visible_variant_ids(product)
    h = headers("shop", 10)
    client.put(f"/v1/cart/items/{mango}", json={"qty": 2}, headers=h)
    client.put(f"/v1/cart/items/{cola}", json={"qty": 1}, headers=h)
    owner = ctx["owner"]

    carts = client.get("/v1/admin/carts", headers=owner).json()
    assert len(carts) == 1 and carts[0]["positions"] == 2 and carts[0]["total"] == 3 * 26000 and carts[0]["can_remind"]
    detail = client.get(f"/v1/admin/carts/{carts[0]['cart_id']}", headers=owner).json()
    assert detail["cart"]["total_qty"] == 3

    r = client.post(f"/v1/admin/carts/{carts[0]['cart_id']}/remind", headers=owner)
    assert r.status_code == 200 and not r.json()["can_remind"]
    reminder = notifications(session_factory, m.NotificationType.cart_reminder)
    assert len(reminder) == 1 and "2 позиции" in reminder[0].text
    assert reminder[0].reply_markup["inline_keyboard"][0][0]["web_app"]["url"] == "https://opt.example.com/t/shop/cart"
    assert client.post(f"/v1/admin/carts/{carts[0]['cart_id']}/remind", headers=owner).status_code == 429

    client.post("/v1/cart/submit", json={}, headers=h)
    assert client.get("/v1/admin/carts", headers=owner).json() == []


def _png() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (3000, 1000), (200, 10, 10)).save(buffer, format="PNG")
    return buffer.getvalue()


def test_photos_upload_resize_and_limit(client, session_factory):
    ctx = setup_shop(client, session_factory, age_gate=False)
    product = create_product(client, ctx)
    url = f"/v1/admin/products/{product['id']}/photos"
    for _ in range(5):
        r = client.post(url, files={"file": ("a.png", _png(), "image/png")}, headers=ctx["owner"])
        assert r.status_code == 201, r.text
    assert client.post(url, files={"file": ("a.png", _png(), "image/png")}, headers=ctx["owner"]).status_code == 409
    assert client.post(f"/v1/admin/products/{product['id']}/photos", files={"file": ("x.txt", b"nope", "text/plain")},
                       headers=ctx["owner"]).status_code in (409, 422)

    photos = r.json()["photos"]
    served = client.get(photos[0]["url"])
    assert served.status_code == 200 and served.headers["content-type"] == "image/webp"
    assert max(Image.open(io.BytesIO(served.content)).size) == 1600

    reordered = [p["id"] for p in reversed(photos)]
    r = client.put(f"{url}/order", json=reordered, headers=ctx["owner"])
    assert [p["id"] for p in r.json()["photos"]] == reordered
    card = client.get(f"/v1/catalog/products/{product['id']}", headers=headers("shop", 10)).json()
    assert card["photos"][0] == r.json()["photos"][0]["url"]
    r = client.delete(f"{url}/{reordered[0]}", headers=ctx["owner"])
    assert len(r.json()["photos"]) == 4
    assert client.get("/v1/files/../../etc/passwd").status_code == 404


def test_ask_manager_link(client, session_factory):
    ctx = setup_shop(client, session_factory, age_gate=False, manager_username="amigo_opt")
    product = create_product(client, ctx)
    card = client.get(f"/v1/catalog/products/{product['id']}", headers=headers("shop", 10)).json()
    assert card["ask_manager_url"].startswith("https://t.me/amigo_opt?text=")


def test_seed_demo_cli(client, session_factory):
    from app.cli import _with_repo, seed_demo

    run(_with_repo(seed_demo, Namespace(slug="demo", owner_telegram_id=555, bot_token=None, bot_username=None)))
    h = headers("demo", 10)
    client.post("/v1/me/age-confirm", headers=h)
    assert client.get("/v1/catalog/products", headers=h).json()["total"] == 6
    assert client.get("/v1/me", headers=headers("demo", 555)).json()["role"] == "owner"


def test_duplicates_are_409(client, session_factory):
    ctx = setup_shop(client, session_factory)
    owner = ctx["owner"]
    assert client.post("/v1/admin/price-tiers", json={"label": "дубль", "min_qty": 10}, headers=owner).status_code == 409
    assert client.post("/v1/admin/brands", json={"name": "Smoant"}, headers=owner).status_code == 201
    assert client.post("/v1/admin/brands", json={"name": "Smoant"}, headers=owner).status_code == 409
    other = client.post("/v1/admin/brands", json={"name": "Elf"}, headers=owner).json()
    # переименование в существующее ловится на commit
    r = client.patch(f"/v1/admin/brands/{other['id']}", json={"name": "Smoant"}, headers=owner)
    assert r.status_code == 409
    assert client.get("/v1/admin/brands", headers=owner).status_code == 200
