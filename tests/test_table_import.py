"""Импорт табличной выгрузки (МойСклад, формат прайса «Галактики»)"""
import io
from argparse import Namespace

from openpyxl import Workbook

import app.database.models as m
from app.importer.detect import parse_price_file
from tests.helpers import headers, run, seed_member, seed_tenant


def build_table() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append([None, "Тип", "Внешний код", "Код товара модификации", "Наименование", "Количество", None, "Цена"])
    ws.append([None, None, None, None, "Галактика", "Связаться с нами"])
    ws.append([None, None, None, None, "Официальный импортер", "Телеграм канал"])
    ws["F3"].hyperlink = "https://t.me/VapeBar_Galaxy"
    ws.append([None, None, None, None, "\n- Минимальная сумма заказа — 20 000 руб\n- Скидка 2% при полной предоплате"])
    ws.append([None, None, None, None, "Наименование", "Кол-во", "Ед.изм.", "От 20 тыс", "От 50 тыс", "От 1 млн", "Сумма"])
    rows = [
        ("", "", "", "Товары/1. ЖЕЛЕЗО (Наборы)/1. GEEKVAPE/Hero 5"),
        ("Модификация", "a1", "01731", "Набор Geekvape Hero 5 Kit (Blaze Red)", 1872, 1811, 1765),
        ("Модификация", "a2", "01731", "Набор Geekvape Hero 5 Kit (Lightning Yellow (Новинка 06.26))", 1872, 1811, 1765),
        ("Модификация", "a3", "01731", "Набор Geekvape Hero 5 Kit (Gold Edition)", 1990, 1950, 1900),
        ("Товар", "a4", "", "Промо Набор Geekvape (5+1)", 0, 0, 0),
        ("", "", "", "Товары/1. ЖЕЛЕЗО (Наборы)/2. VAPORESSO/VIBE - Горячее предложение!"),
        ("Модификация", "b1", "01704", "Набор Vaporesso VIBE 1100mAh 24W KIT (Black)", 599, 599, 599),
        ("", "", "", "Товары/2. РАСХОДНИКИ/ИСПАРИТЕЛИ/VAPORESSO"),
        ("Товар", "c1", "", "Испаритель Vaporesso GTX 0,2 Ом Mesh Coil (5шт)", 590, 580, 555),
        ("Товар", "c2", "", "Испаритель Vaporesso GTX 1,2 Ом Mesh Coil (5шт)", 600, 590, 560),
        ("Товар", "c3", "", "Обслуживаемая база Smoant K-RBA", 290, 277, 263),
        ("", "", "", "Товары/4. ПРОМО (МЕРЧ)/Vaporesso"),
        ("Товар", "d1", "", "Vaporesso Сумка поясная", 0, 0, 0),
    ]
    for t, ext, code, name, *prices in rows:
        ws.append([None, t or None, ext or None, code or None, name, None, "шт" if t else None, *prices])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_table_parser():
    kind, res = parse_price_file(build_table(), "galaxy.xlsx")
    assert kind == "table"
    assert res.tenant_defaults["min_order_amount"] == 20000
    assert res.tenant_defaults["manager_username"] == "VapeBar_Galaxy"
    products = {p.name: p for p in res.products}
    hero = products["Набор Geekvape Hero 5 Kit"]
    assert hero.brand == "GEEKVAPE" and hero.sku == "01731" and hero.category_path == ["Железо (наборы)"]
    assert hero.prices == {20000: 187200, 50000: 181100, 1000000: 176500}
    assert [(v.name, v.prices) for v in hero.variants] == [
        ("Blaze Red", None), ("Lightning Yellow", None), ("Gold Edition", {20000: 199000, 50000: 195000, 1000000: 190000}),
    ]
    assert products["Набор Vaporesso VIBE 1100mAh 24W KIT"].description == "Горячее предложение"
    gtx = products["Испаритель Vaporesso GTX Mesh Coil (5шт)"]
    assert gtx.category_path == ["Расходники", "Испарители"] and [v.name for v in gtx.variants] == ["0,2 Ом", "1,2 Ом"]
    assert "Обслуживаемая база Smoant K-RBA" in products
    assert not any("Промо" in n or "Сумка" in n for n in products)  # нулевые цены и «Промо» пропущены


def test_table_import(client, session_factory, tmp_path):
    from app.cli import _with_repo, import_price

    tenant_id = seed_tenant(session_factory, "galactica", age_gate=False)
    seed_member(session_factory, tenant_id, 1, role=m.Role.owner)
    owner = headers("galactica", 1)
    path = tmp_path / "galaxy.xlsx"
    path.write_bytes(build_table())
    run(_with_repo(import_price, Namespace(slug="galactica", file=str(path), wipe=False, no_photos=True, dry_run=False, no_enrich=True)))

    tenant = client.get("/v1/me", headers=owner).json()["tenant"]
    assert tenant["price_basis"] == "amount" and tenant["min_order_amount"] == 2000000
    assert tenant["manager_username"] == "VapeBar_Galaxy"
    tree = client.get("/v1/catalog/categories", headers=owner).json()
    assert [(c["name"], [ch["name"] for ch in c["children"]]) for c in tree] == [
        ("Железо (наборы)", []), ("Расходники", ["Испарители"]),
    ]
    products = {p["name"]: p for p in client.get("/v1/admin/products", headers=owner).json()["items"]}
    hero = client.get(f"/v1/catalog/products/{products['Набор Geekvape Hero 5 Kit']['id']}", headers=owner).json()
    assert hero["sku"] == "01731"
    gold = next(v for v in hero["variants"] if v["name"] == "Gold Edition")
    assert [p["amount"] for p in gold["prices"]] == [199000, 195000, 190000]

    # клиент: заявка ниже минимальной 20 000 ₽ не уходит
    h = headers("galactica", 10)
    cart = client.put(f"/v1/cart/items/{gold['id']}", json={"qty": 2}, headers=h).json()
    assert cart["blockers"] == ["below_min_amount"] and cart["total"] == 2 * 199000


def test_named_tiers_and_top_level_groups():
    """1С «Остатки»: «Мелкий ОПТ» / «Крупный ОПТ» без порогов, группы без «/», кириллические бренды"""
    import openpyxl
    from io import BytesIO
    from app.importer.table_price import parse_table

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Остатки Просто ОПТ"])
    ws.append(["Наименование", "Мелкий ОПТ", "Крупный ОПТ", "Заказ"])
    for brand in ("Angry Ape", "BJORN", "DUALL", "SKALA", "OGGO", "Злая Монашка"):
        ws.append([f"Жидкости/{brand}"])
        ws.append([f"{brand} Hard (Вишня)", 265, 255])
        ws.append([f"{brand} Hard (Манго)", 265, 255])
    ws.append(["Поды и расходники от 40т"])
    ws.append(["Аккумулятор LG 18650 HG2", 300, 280])
    buf = BytesIO()
    wb.save(buf)

    result = parse_table(buf.getvalue(), "ostatki.xlsx")
    assert result.tenant_defaults["tier_labels"] == {1: "Мелкий опт", 40000: "Крупный опт от 40 000 ₽"}
    assert "min_order_amount" not in result.tenant_defaults
    by_name = {p.name: p for p in result.products}
    monashka = by_name["Злая Монашка Hard"]
    assert monashka.brand == "Злая Монашка" and monashka.category_path == ["Жидкости"]
    assert monashka.prices == {1: 26500, 40000: 25500} and [v.name for v in monashka.variants] == ["Вишня", "Манго"]
    assert by_name["Аккумулятор LG 18650 HG2"].category_path == ["Поды и расходники"]
