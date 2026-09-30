"""
Парсер «ручных» прайсов-витрин (формат «Сафари вейп», Google Sheets):
- лист = категория; строка заголовков «Название | Вкусы/Расцветки/Мощность |
  (Крепость) | Цена» (у маленьких листов её может не быть — тогда B/C/D);
- товар = блок строк: фото в объединённой ячейке колонки A (или название в
  объединённой B), вкусы/цвета — по одному в строке колонки C;
- строка, объединённая на всю ширину, — бренд («Vaporesso (к каждому
  устройству жижа в подарок🎁)» → бренд Vaporesso, остальное — в описание);
- красная заливка ячейки вкуса (или названия у товара без вкусов) — нет в наличии;
- «(ПРИ ПОКУПКЕ ОТ 2 шт / Цена будет 550₽)» под ценой — второй уровень по количеству.
Чистый модуль — без БД и FastAPI.
"""
from __future__ import annotations

import re
from io import BytesIO

import openpyxl

from app.importer.block_price import ParsedProduct, ParsedSheet, ParsedVariant, ParseResult

NAME_HEADERS = ("название",)
PRICE_HEADERS = ("цена",)
STRENGTH_HEADERS = ("крепость",)
RED = {"FFFF0000", "FFEA4335", "FFE06666", "FFCC0000"}
PRICE_RE = re.compile(r"^\s*(\d[\d\s]*(?:[.,]\d+)?)\s*(?:₽|р\.?|руб\.?)?\s*$", re.I)
BULK_QTY_RE = re.compile(r"от\s*(\d+)\s*шт", re.I)
BULK_PRICE_RE = re.compile(r"цена\s+будет\s*(\d[\d\s]*)", re.I)
NEW_MARK_RE = re.compile(r"^(new|новинка!*|новая линейка)$", re.I)
DANGLING_RE = re.compile(r"\b(на|для|с|и)$", re.I)
FITS_RE = re.compile(r",?\s*подходит\s+на\b.*$", re.I)
EMOJI_RE = re.compile("[\U0001F000-\U0001FAFF\u2600-\u27BF\uFE0F]")
STRENGTH_ONLY_RE = re.compile(r"^\d+\s*(мг|mg)$", re.I)


def _text(v) -> str:
    return "" if v is None else re.sub(r"[ \t]+", " ", str(v)).strip()


def _price(v) -> int | None:
    """1450.0 / «2600₽» / «350р» -> копейки"""
    if isinstance(v, (int, float)):
        return round(v * 100) if v > 0 else None
    if m := PRICE_RE.match(_text(v)):
        number = float(re.sub(r"\s", "", m.group(1)).replace(",", "."))
        return round(number * 100) if number > 0 else None
    return None


def _red(cell) -> bool:
    try:
        return cell.fill.fill_type == "solid" and str(cell.fill.fgColor.rgb).upper() in RED
    except Exception:
        return False


def _header(ws) -> dict | None:
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=20), start=1):
        texts = [_text(c.value).lower() for c in row]
        name = next((i for i, t in enumerate(texts) if t.startswith(NAME_HEADERS)), None)
        price = next((i for i, t in enumerate(texts) if t.startswith(PRICE_HEADERS)), None)
        if name is None or price is None:
            continue
        strength = next((i for i, t in enumerate(texts) if t.startswith(STRENGTH_HEADERS)), None)
        return {"row": r, "name": name, "variant": name + 1, "strength": strength, "price": price}
    return None


def _vertical_starts(ws, first_row: int, width: int) -> tuple[set[int], dict[int, int]]:
    """Начала товарных блоков и строки-бренды (объединения на всю ширину).
    Блок — объединённая ячейка фото в A; объединение в B считается началом, только если
    не лежит внутри блока A (иначе это просто многострочное название внутри карточки)."""
    a_ranges: list[tuple[int, int]] = []
    b_starts: list[int] = []
    banners: dict[int, int] = {}
    for rng in ws.merged_cells.ranges:
        if rng.min_row < first_row:
            continue
        if rng.min_col == 1 and rng.max_col >= width:
            banners[rng.min_row] = rng.max_row
        elif rng.min_col == rng.max_col == 1 and rng.max_row > rng.min_row:
            a_ranges.append((rng.min_row, rng.max_row))
        elif rng.min_col == rng.max_col == 2 and rng.max_row > rng.min_row:
            b_starts.append(rng.min_row)
    starts = {s for s, _ in a_ranges}
    starts |= {b for b in b_starts if not any(s < b <= e for s, e in a_ranges)}
    # цена в строке без объединённого фото — отдельная карточка («Испаритель P-coil» под «s-coil»)
    for r in range(first_row, ws.max_row + 1):
        if any(s <= r <= e for s, e in a_ranges) or r in banners:
            continue
        if _price(ws.cell(r, width).value) and any(_text(ws.cell(r, c).value) for c in (2, 3)):
            starts.add(r)
    return starts, banners


def is_merged_price(content: bytes) -> bool:
    try:
        wb = openpyxl.load_workbook(BytesIO(content))
    except Exception:
        return False
    for ws in wb.worksheets[:3]:
        header = _header(ws)
        if header and len(_vertical_starts(ws, header["row"] + 1, header["price"] + 1)[0]) >= 3:
            return True
    return False


DEFAULT_ROW_EMU = 190500  # 15pt


def _image_center_row(ws, img) -> int:
    """Строка, на которую приходится середина картинки (якорь часто стоит на строку-две выше блока)"""
    anchor = img.anchor
    row = anchor._from.row + 1
    if hasattr(anchor, "to") and anchor.to is not None:
        return (anchor._from.row + anchor.to.row) // 2 + 1
    remaining = anchor._from.rowOff + anchor.ext.height / 2
    while True:
        height = ws.row_dimensions[row].height
        step = height * 12700 if height else DEFAULT_ROW_EMU
        if remaining < step:
            return row
        remaining -= step
        row += 1


def _sheet_title(title: str) -> str:
    """«СнюсПластинки» -> «Снюс, пластинки», «Под Устройства» -> «Под-устройства»"""
    title = re.sub(r"(?<=[а-яa-z])(?=[А-ЯA-Z])", ", ", title.strip())
    words = title.split(" ")
    title = " ".join([words[0]] + [w.lower() for w in words[1:]])
    if title.lower().startswith("под "):
        title = "Под-" + title[4:]
    return title


def _banner(text: str) -> tuple[str, str | None]:
    """«Vaporesso (к каждому устройству жижа в подарок🎁)» -> («Vaporesso», «к каждому …»)"""
    text = text.strip(" .")
    note = None
    if m := re.search(r"\((.+)\)\s*$", text):
        note = m.group(1).strip()
        text = text[: m.start()].strip(" .")
    if text.isupper() and len(text) > 3:
        text = text.title()
    return text, note


def _block_name(lines: list[str]) -> tuple[str, list[str], bool]:
    """Название — первая содержательная строка (пометки «NEW» — в is_new), остальное — описание"""
    parts = [p.strip() for line in lines for p in line.split("\n") if p.strip()]
    is_new = False
    while parts and NEW_MARK_RE.match(parts[0]):
        is_new = True
        parts.pop(0)
    if not parts:
        return "", [], is_new
    name = parts.pop(0)
    rest: list[str] = []
    if m := FITS_RE.search(name):
        name = name[: m.start()].strip(" ,")
        devices = [p for p in parts if not p.lower().startswith("устройства")]
        rest.append("Подходит на: " + " ".join(devices).strip(" ,."))
        parts = []
    while parts and DANGLING_RE.search(name):
        name = f"{name} {parts.pop(0)}"
    name = _clean(name)
    rest += [p.strip(" ~") for p in parts]
    return name, rest, is_new


def _clean(text: str) -> str:
    """Эмодзи и непарные скобки: «Никотин. ЖВАЧКИ 😋)» -> «Никотин. ЖВАЧКИ»"""
    text = EMOJI_RE.sub("", text).strip()
    if text.count(")") > text.count("("):
        text = text.rstrip(") ")
    return re.sub(r"\s+", " ", text).strip()


def parse_merged(content: bytes, with_images: bool = True) -> ParseResult:
    wb = openpyxl.load_workbook(BytesIO(content))
    sheets: list[ParsedSheet] = []
    warnings: list[str] = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        header = _header(ws) or {"row": 0, "name": 1, "variant": 2, "strength": None, "price": 3}
        first = header["row"] + 1
        width = header["price"] + 1
        starts, banners = _vertical_starts(ws, first, width)
        # строки-бренды и шапка — не товары; продолжение шапки (объединение над таблицей) пропускаем
        banner_rows = {r for s, e in banners.items() for r in range(s, e + 1)}
        starts = {s for s in starts if s not in banner_rows}
        # фото (A) и название (B) бывают объединены со сдвигом на строку — это один блок
        merged_starts: list[int] = []
        for r in sorted(starts):
            if merged_starts and r - merged_starts[-1] <= 2:
                continue
            merged_starts.append(r)
        starts = set(merged_starts)
        bounds = sorted(starts | set(banners))
        images: list[tuple[int, bytes]] = []
        if with_images:
            for img in ws._images:
                try:
                    images.append((_image_center_row(ws, img), img._data()))
                except Exception:
                    continue

        title = _sheet_title(ws.title)
        brand: str | None = None
        brand_note: str | None = None
        products: list[ParsedProduct] = []
        for idx, start in enumerate(bounds):
            end = (bounds[idx + 1] - 1) if idx + 1 < len(bounds) else ws.max_row
            if start in banners:
                brand, brand_note = _banner(_text(ws.cell(start, 1).value))
                continue
            products.extend(_block(ws, header, start, end, title, brand, brand_note, images, warnings))
        if products:
            sheets.append(ParsedSheet(title=title, products=products))
    return ParseResult(sheets=sheets, warnings=warnings, tenant_defaults={"price_basis": "qty"})


def _block(ws, header, start, end, title, brand, brand_note, images, warnings) -> list[ParsedProduct]:
    name_col, var_col, price_col = header["name"] + 1, header["variant"] + 1, header["price"] + 1
    names: list[tuple[int, str]] = []
    variants: list[tuple[int, str, bool]] = []
    price = None
    bulk_qty = bulk_price = None
    strength = None
    name_red = False
    for r in range(start, end + 1):
        name_cell = ws.cell(r, name_col)
        if text := _text(name_cell.value):
            names.append((r, text))
            name_red = name_red or _red(name_cell)
        var_cell = ws.cell(r, var_col)
        if (v := _text(var_cell.value)) and len(v) > 1:
            variants.append((r, v, _red(var_cell)))
        raw_price = ws.cell(r, price_col).value
        if price is None and (p := _price(raw_price)):
            price = p
        elif isinstance(raw_price, str):
            if m := BULK_QTY_RE.search(raw_price):
                bulk_qty = int(m.group(1))
            if m := BULK_PRICE_RE.search(raw_price):
                bulk_price = int(re.sub(r"\s", "", m.group(1))) * 100
        if header["strength"] is not None and strength is None:
            raw = ws.cell(r, header["strength"] + 1).value
            strength = (f"{raw:g} мг" if isinstance(raw, (int, float)) else _text(raw)) or None
    if not names and not variants:
        return []
    if price is None:
        warnings.append(f"{ws.title}, строки {start}–{end}: «{names[0][1] if names else variants[0][1]}» — нет цены, пропущен")
        return []
    prices = {1: price}
    if bulk_qty and bulk_price and bulk_price < price:
        prices[bulk_qty] = bulk_price

    image = None
    for row, data in images:
        if start <= row <= end:
            image = data
            break

    # «Так себе 150мг / 80мг / 40мг» в одном блоке — отдельные товары, каждый со своими вкусами
    groups: list[tuple[list[str], list[tuple[int, str, bool]]]] = []
    split_rows = [r for r, t in names if any(vr == r for vr, _, _ in variants)]
    if len(split_rows) >= 2 and _same_family([t for r, t in names if r in split_rows]):
        bounds = split_rows + [end + 1]
        for i, r in enumerate(split_rows):
            groups.append((
                [t for nr, t in names if nr == r],
                [v for v in variants if r <= v[0] < bounds[i + 1]],
            ))
    elif names and all(STRENGTH_ONLY_RE.match(t.strip()) for _, t in names) and variants:
        # «150мг / 70мг» без марки, вкус «Original»: товар — вкус, варианты — крепость
        groups.append(([variants[0][1]], [(r, t, _red(ws.cell(r, name_col))) for r, t in names]))
    else:
        groups.append(([t for _, t in names], variants))

    result = []
    for lines, block_variants in groups:
        name, rest, is_new = _block_name(lines)
        if not name:
            name = block_variants[0][1] if block_variants else ""
            block_variants = block_variants[1:] if block_variants else []
        if not name:
            continue
        if brand and STRENGTH_ONLY_RE.match(name):
            name = f"{brand} {name}"
        notes = []
        if is_new:
            notes.append("Новинка")
        # «В комплекте; 1 картридж» в колонке вкусов — это описание, а не варианты
        if block_variants and any(v.endswith((";", ":")) for _, v, _ in block_variants):
            rest.append(" ".join(v for _, v, _ in block_variants))
            block_variants = []
        if rest:
            notes.append(re.sub(r"\s+", " ", " ".join(rest)).strip())
        if strength and not re.search(r"\d\s*(мг|mg)", name, re.I):
            notes.append(f"Крепость: {strength}")
        if len(prices) > 1:
            (qty, bulk), = [(q, p) for q, p in prices.items() if q != 1]
            notes.append(f"От {qty} шт — {bulk // 100} ₽ за штуку")
        if brand_note:
            notes.append(brand_note)
        seen: set[str] = set()
        parsed_variants = []
        for _, v, red in block_variants:
            v = v.strip(" ,;.")
            if re.fullmatch(r"\d+(?:[.,]\d+)?", v):
                v = f"{v.replace('.', ',')} Ом"
            if v.lower() in seen:
                continue
            seen.add(v.lower())
            parsed_variants.append(ParsedVariant(name=v, out=red))
        out = all(v.out for v in parsed_variants) if parsed_variants else name_red
        result.append(ParsedProduct(
            sheet=title, row=start, name=name, brand=brand, prices=dict(prices), variants=parsed_variants,
            description=" · ".join(notes) or None, out=out, image=image,
        ))
    return result


def _same_family(names: list[str]) -> bool:
    first = [n.split()[0].lower() for n in names if n.split()]
    return len(set(first)) == 1 or all(STRENGTH_ONLY_RE.match(n.strip()) for n in names)
