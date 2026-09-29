"""
Парсер «блочных» прайсов (формат Amigo Opt): не таблица, а карточки товаров.
Чистый модуль — без БД и FastAPI; результат применяет
app/services/price_import.py.

Как устроен файл:
- лист = категория (скрытые листы и «О НАС» пропускаются), строки 1–13 —
  шапка компании;
- заголовок товара — ячейка A с «товарной» заливкой (оранжевая/голубая),
  зачёркнутый заголовок = нет в наличии;
- бренд — ячейка A с «брендовой» заливкой («ПРОДУКЦИЯ ICON»), а на листах
  без таких заливок (жидкости) — одиночная строка прямо перед товаром;
- цены: столбиком («ОТ 3.000» → цена в соседней справа ячейке, «РОЗНИЦА»)
  или строкой (подписи уровней в строке заголовка, цены — в следующей);
- вкусы/цвета: всё остальное текстовое в блоке — строки под товаром,
  колонки, многострочные ячейки со списками «1. …», «1- …», «1) …»;
  зачёркнутый вкус = нет в наличии;
- описание: «20 вкусов», «МИНИМАЛЬНЫЙ ЗАКАЗ…», характеристики с «:»,
  «N ШТ В УПАКОВКЕ»;
- «СТРОГО МИКС» — вкус не выбирается, товар без вариантов;
- картинка, привязанная к строкам блока, — фото товара.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from io import BytesIO

import openpyxl

PRODUCT_FILLS = {"FF9900", "E69138", "6FA8DC"}
BRAND_FILLS = {"F9CB9C", "F6B26B", "3D85C6"}
HEADER_ROWS = 13  # шапка компании на каждом листе
SKIP_SHEETS = {"О НАС!", "О НАС"}

# Служебные надписи внутри блока — не вкусы
JUNK = (
    "ГРАДАЦИЯ ЦЕН", "КРУПНЫЙ ОБЪЁМ", "КРУПНЫЙ ОБЪЕМ", "НОВИНКА", "РАСЦВЕТКИ", "ЦВЕТА", "ЦЕНА",
    "ФОТОГРАФИЮ ПОЗИЦИИ", "ДОГОВОРИМСЯ", "ИНДИВИДУАЛЬНОЕ ПРЕДЛОЖЕНИЕ",
)
MIX_MARKS = ("СТРОГО МИКС", "ТОЛЬКО МИКС")
# «ОТ 3.000», «ОТ 30.000Р», «3.000», «ОТ 3000» — но не голое число «260» (это цена)
TIER_LABEL_RE = re.compile(r"^(?:ОТ\s*(\d+(?:[.\s]\d{3})*)|(\d{1,3}(?:[.\s]\d{3})+))\s*Р?\.?$", re.IGNORECASE)
# «РОЗНИЦА ОТ 1.000» и «Любая сумма» (одна цена на все уровни) — младший уровень
RETAIL_RE = re.compile(r"РОЗНИЦА|ЛЮБАЯ\s+СУММА", re.IGNORECASE)
RETAIL_THRESHOLD = 1000  # «РОЗНИЦА ОТ 1.000»
# «1.», «1)», «1-», «1 -», «1➖», «1 Бабл гам», «◦», «- », «•» в начале строки списка вкусов
# (цифра без разделителя — часть названия: «7UP лимон»)
NUMBERING_RE = re.compile(r"^\s*(?:\d{1,3}(?:[️⃣]+|\s*[.),\-–—:➖](?!\d)|\s+(?=\D))|[◦•➖▪️\-–—*️]+)[\s⁠️]*")
# «Вишня \ 13. Зелёный виноград» — несколько вкусов в строке через обратный слэш
BACKSLASH_SPLIT_RE = re.compile(r"\s*\\\s*")
DESCRIPTION_RE = re.compile(
    r"(\d+\s*вкус|МИНИМАЛЬН|В\s+УПАКОВК|\bШТ\s+В\b|ГАРАНТИ|ИДЕНТИФИЦ|ИДЕНФИЦ|IMEI|ЗАРЯДК|АККУМУЛЯТОР|ВРЕМЯ\s+РАБОТ|\bДЛЯ\s|ВАННОЧКА|ШУМОПОДАВЛ|\bIOS\b|\bЧИП\b)",
    re.IGNORECASE,
)
RESISTANCE_SPLIT_RE = re.compile(r"\s{3,}")
PACKAGING_RE = re.compile(r"\d+\s*ШТ\s+В\s+УПАКОВКЕ", re.IGNORECASE)


@dataclass
class ParsedVariant:
    name: str
    out: bool = False


@dataclass
class ParsedProduct:
    sheet: str
    row: int
    name: str
    brand: str | None
    prices: dict[int, int]  # порог уровня в рублях (1000 — розница) -> цена в копейках
    variants: list[ParsedVariant] = field(default_factory=list)
    description: str | None = None
    out: bool = False
    mix: bool = False
    image: bytes | None = None


@dataclass
class ParsedSheet:
    title: str
    products: list[ParsedProduct]


@dataclass
class ParseResult:
    sheets: list[ParsedSheet]
    warnings: list[str]

    @property
    def products(self) -> list[ParsedProduct]:
        return [p for s in self.sheets for p in s.products]


def _fill(cell) -> str:
    try:
        if cell.fill and cell.fill.fill_type == "solid":
            rgb = cell.fill.fgColor.rgb
            return rgb[-6:].upper() if isinstance(rgb, str) else ""
    except Exception:
        pass
    return ""


def _strike(cell) -> bool:
    return bool(cell.font and cell.font.strike)


def _text(value) -> str:
    return str(value).replace("⁠", " ").strip() if value is not None else ""


def _clean_spaces(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _number(value) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    text = _text(value).replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def tier_threshold(label: str) -> int | None:
    """«ОТ 3.000» / «ОТ 30.000Р» / «10.000» -> 3000 / 30000 / 10000; «РОЗНИЦА…» -> 1000"""
    text = _clean_spaces(label).upper()
    if RETAIL_RE.search(text):
        return RETAIL_THRESHOLD
    m = TIER_LABEL_RE.match(text)
    if not m:
        return None
    return int(re.sub(r"[.\s]", "", m.group(1) or m.group(2)))


def _is_junk(text: str) -> bool:
    upper = text.upper()
    return any(mark in upper for mark in JUNK)


def clean_variant_name(line: str) -> str:
    """«1.⁠ ⁠Мятная Жвачка айс» / «3 - Белый виноград» / «4) яблочная вишня» / «1- табак яблоко мята)» -> название"""
    text = NUMBERING_RE.sub("", line.replace("⁠", " "))
    text = _clean_spaces(text).strip(" .;,")
    if text.endswith(")") and "(" not in text:
        text = text[:-1].strip()
    if (text.startswith("[") and text.endswith("]")) or (text.startswith("(") and text.endswith(")")):
        text = text[1:-1].strip()
    return text


def _is_description(text: str) -> bool:
    lines = [l for l in text.splitlines() if l.strip()]
    if not lines:
        return False
    if DESCRIPTION_RE.search(text):
        return True
    # «1: Арбуз вишня» — нумерация списка вкусов, а не «Характеристика: значение»
    with_colon = sum(1 for l in lines if ":" in NUMBERING_RE.sub("", l, count=1) and not re.match(r"^\s*\d{1,3}\s*:", l))
    return with_colon * 2 >= len(lines)


def _split_variants(text: str) -> list[str]:
    names = []
    lines = [part for line in text.splitlines() for part in BACKSLASH_SPLIT_RE.split(line)]
    for line in lines:
        # «Coil 0.6 om       Coil 1.2 om» — несколько вариантов в строке через длинные пробелы
        parts = RESISTANCE_SPLIT_RE.split(line) if re.search(r"(om|ом)\b", line, re.IGNORECASE) else [line]
        for part in parts:
            name = clean_variant_name(part)
            if name and not _is_junk(name) and len(name) <= 120:
                names.append(name)
    return names


def _product_name(raw: str) -> tuple[str, str | None]:
    """Многострочный заголовок «КАРТРИДЖ\\nCHARON T50\\n1 ШТ В УПАКОВКЕ» -> («КАРТРИДЖ CHARON T50», «1 ШТ В УПАКОВКЕ»)"""
    lines = [_clean_spaces(l) for l in raw.replace("⁠", " ").splitlines() if l.strip()]
    packaging = [l for l in lines if PACKAGING_RE.search(l)]
    name_lines = [PACKAGING_RE.sub("", l).strip() for l in lines]
    name = _clean_spaces(" ".join(l for l in name_lines if l).replace("()", ""))
    note = _clean_spaces(" ".join(PACKAGING_RE.search(l).group(0) for l in packaging)) if packaging else None
    return name, note


def _brand_name(raw: str) -> str | None:
    text = _clean_spaces(raw)
    text = re.sub(r"^(ПРОДУКЦИЯ|ПРОИЗВОДИТЕЛЬ)\s+", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s*\(ОРИГИНАЛ\)\s*$", "", text, flags=re.IGNORECASE).strip()
    # Длинные заголовки групп («Часы с круглым дисплеем», «СМАРТ ЧАСЫ на приложении…») — не бренды
    return text if text and len(text.split()) <= 3 else None


def _brand_in_name(brand: str, name: str) -> bool:
    def norm(text: str) -> str:
        return re.sub(r"[^a-zа-я0-9]", "", text.lower().replace("ё", "е"))
    token = norm(brand.split()[0]) if brand.split() else ""
    return bool(token) and token in norm(name)


def parse_workbook(content: bytes, with_images: bool = True, include_hidden: bool = False) -> ParseResult:
    wb = openpyxl.load_workbook(BytesIO(content))
    sheets: list[ParsedSheet] = []
    warnings: list[str] = []
    for ws in wb.worksheets:
        if ws.title.strip() in SKIP_SHEETS or (ws.sheet_state != "visible" and not include_hidden):
            continue
        parsed = _parse_sheet(ws, with_images, warnings)
        if parsed.products:
            sheets.append(parsed)
    return ParseResult(sheets=sheets, warnings=warnings)


def _parse_sheet(ws, with_images: bool, warnings: list[str]) -> ParsedSheet:
    rows: list[tuple[int, list]] = []
    for r, row in enumerate(ws.iter_rows(min_row=HEADER_ROWS + 1), start=HEADER_ROWS + 1):
        cells = [c for c in row if _text(c.value)]
        if cells:
            rows.append((r, cells))

    def is_header(cells) -> bool:
        first = cells[0]
        return first.column == 1 and _fill(first) in PRODUCT_FILLS

    has_brand_fills = any(cells[0].column == 1 and _fill(cells[0]) in BRAND_FILLS for _, cells in rows)
    header_idx = [i for i, (_, cells) in enumerate(rows) if is_header(cells)]

    images: dict[int, bytes] = {}
    if with_images:
        for img in ws._images:
            try:
                images.setdefault(img.anchor._from.row + 1, img._data())
            except Exception:
                continue

    products: list[ParsedProduct] = []
    brand: str | None = None
    brand_rows: set[int] = set()
    # Бренды: заливка «ПРОДУКЦИЯ …», на листах без неё — одиночная строка перед товаром
    brand_before: dict[int, str] = {}
    for i, (r, cells) in enumerate(rows):
        first = cells[0]
        if first.column == 1 and _fill(first) in BRAND_FILLS and not is_header(cells):
            brand_rows.add(r)
    if not has_brand_fills:
        # Жидкости: бренд — одиночная строка прямо перед товаром (через одну пустую),
        # отделённая пустой строкой и от предыдущего текста; хвост списка вкусов
        # предыдущего товара стоит вплотную к нему и брендом не считается.
        for hi in header_idx:
            if hi == 0 or hi - 1 in header_idx:
                continue
            r, cells = rows[hi - 1]
            text = _text(cells[0].value)
            if len(cells) != 1 or cells[0].column != 1 or "\n" in text or tier_threshold(text):
                continue
            prefixed = re.match(r"^(ПРОДУКЦИЯ|ПРОИЗВОДИТЕЛЬ)\s", text, re.IGNORECASE)
            gap_before = r - rows[hi - 2][0] if hi >= 2 else 99
            gap_after = rows[hi][0] - r
            if prefixed or (gap_before >= 2 and gap_after <= 3 and not NUMBERING_RE.match(text)):
                brand_rows.add(r)
                brand_before[rows[hi][0]] = _brand_name(text)

    for n, hi in enumerate(header_idx):
        header_row, header_cells = rows[hi]
        # бренд: ближайшая брендовая строка выше, не дальше предыдущего товара
        start_scan = header_idx[n - 1] + 1 if n else 0
        direct = False
        for j in range(hi - 1, start_scan - 1, -1):
            if rows[j][0] in brand_rows:
                brand = brand_before.get(header_row) or _brand_name(_text(rows[j][1][0].value))
                direct = True
                break
        if not direct and brand and not _brand_in_name(brand, _text(header_cells[0].value)):
            # Бренд предыдущего блока переносим, только если он есть в названии товара
            # («KASTA COVID» под «KASTA»), иначе это другой производитель без своего заголовка
            brand = None
        end = header_idx[n + 1] if n + 1 < len(header_idx) else len(rows)
        block = [rows[k] for k in range(hi, end) if rows[k][0] not in brand_rows or k == hi]
        next_row = rows[end][0] if end < len(rows) else ws.max_row + 1
        product = _parse_block(ws.title, header_row, header_cells, block, brand)
        if product is None:
            continue
        for image_row in range(header_row, next_row):
            if image_row in images:
                product.image = images[image_row]
                break
        if not product.prices:
            warnings.append(f"{ws.title}, строка {header_row}: «{product.name}» — не нашёл цен, пропущен")
            continue
        products.append(product)

    # одинаковые названия внутри листа — разводим номером («MFU 25000» дважды с разными ценами)
    seen: dict[str, int] = {}
    for p in products:
        key = p.name.lower()
        seen[key] = seen.get(key, 0) + 1
        if seen[key] > 1:
            warnings.append(f"{ws.title}: повтор «{p.name}» — назван «{p.name} #{seen[key]}»")
            p.name = f"{p.name} #{seen[key]}"
    return ParsedSheet(title=ws.title.strip(), products=products)


def _parse_block(sheet: str, header_row: int, header_cells, block, brand: str | None) -> ParsedProduct | None:
    name_raw = _text(header_cells[0].value)
    name, packaging = _product_name(name_raw)
    if not name:
        return None
    product = ParsedProduct(sheet=sheet, row=header_row, name=name, brand=brand, prices={}, out=_strike(header_cells[0]))
    descriptions: list[str] = [packaging] if packaging else []
    used: set[tuple[int, int]] = {(header_row, header_cells[0].column)}

    # Горизонтальные уровни: подписи «ОТ 3.000 …» в строке заголовка, цены — в следующей строке
    # (подпись с числом справа в той же строке — это вертикальная раскладка: «Любая сумма | 360»)
    header_by_col = {c.column: c for c in header_cells}
    horizontal = {
        c.column: tier_threshold(_text(c.value)) for c in header_cells[1:]
        if isinstance(c.value, str) and tier_threshold(_text(c.value))
        and _number(getattr(header_by_col.get(c.column + 1), "value", None)) is None
    }
    if horizontal:
        for c in header_cells[1:]:
            used.add((header_row, c.column))
        for r, cells in block[1:3]:
            prices = {c.column: _number(c.value) for c in cells if c.column in horizontal}
            if any(v is not None for v in prices.values()):
                for col, value in prices.items():
                    if value is not None and value > 0:
                        product.prices[horizontal[col]] = round(value * 100)
                for c in cells:
                    if c.column in horizontal:
                        used.add((r, c.column))
                break

    # Вертикальные уровни: подпись -> цена в соседней справа ячейке
    for r, cells in block:
        by_col = {c.column: c for c in cells}
        for c in cells:
            if (r, c.column) in used or not isinstance(c.value, str):
                continue
            threshold = tier_threshold(_text(c.value))
            if threshold is None:
                continue
            price_cell = by_col.get(c.column + 1)
            used.add((r, c.column))
            if price_cell is not None:
                used.add((r, price_cell.column))
                value = _number(price_cell.value)
                if value is not None and value > 0 and threshold not in product.prices:
                    product.prices[threshold] = round(value * 100)

    # Остальной текст блока — вкусы/цвета или описание
    variants: list[ParsedVariant] = []
    for r, cells in block:
        for c in cells:
            if (r, c.column) in used:
                continue
            text = _text(c.value)
            if not text or _number(c.value) is not None:
                continue
            upper = text.upper()
            if any(mark in upper for mark in MIX_MARKS):
                product.mix = True
                continue
            if _is_junk(text) or tier_threshold(text):
                continue
            if _is_description(text):
                descriptions.append(_clean_spaces(text) if "\n" not in text else text.strip())
                continue
            strike = _strike(c)
            for variant_name in _split_variants(text):
                variants.append(ParsedVariant(name=variant_name, out=strike or product.out))

    unique: dict[str, ParsedVariant] = {}
    for v in variants:
        unique.setdefault(v.name.lower().replace("ё", "е"), v)
    product.variants = [] if product.mix else list(unique.values())
    if product.mix:
        descriptions.append("Строго микс — вкусы ассорти, предпочтение можно указать в комментарии к заявке")
    product.description = "\n".join(dict.fromkeys(d for d in descriptions if d)) or None
    return product
