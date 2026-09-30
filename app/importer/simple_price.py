"""
Парсер простых прайсов «строка = товар» (формат «Дымка», Google Sheets):
- лист = категория; в строке заголовков «Наименование товара», колонка
  вкусов/цветов/сопротивлений («Вкусы», «Сопративление»…), уровни цен по
  КОЛИЧЕСТВУ: «от 5 штук» / «от 50 штук», а у расходки — «1 пачка» /
  «Блок(10 пачек)» / «10 блоков»;
- вкусы — списком в одной ячейке (через перенос строки);
- «❌» в названии = нет в наличии;
- картинка в строке товара — фото;
- «1 шт - 720 (Пачка 4 шт)» — цена за штуку и размер пачки: товар продаётся
  пачками, цена пачки = 720 × 4.
Чистый модуль — без БД и FastAPI.
"""
from __future__ import annotations

import re
from io import BytesIO

import openpyxl

from app.importer.block_price import ParsedProduct, ParsedSheet, ParsedVariant, ParseResult

NAME_HEADERS = ("наименование",)
VARIANT_HEADERS = ("вкус", "цвет", "сопротив", "сопратив", "вариант", "модиф")
QTY_TIER_RE = re.compile(r"^от\s*(\d+)\s*(шт|штук)", re.IGNORECASE)
PACK_TIER_RE = re.compile(r"^(\d+)\s*пач", re.IGNORECASE)
BLOCK_TIER_RE = re.compile(r"блок\s*\(?\s*(\d+)\s*пач", re.IGNORECASE)
BLOCKS_RE = re.compile(r"^(\d+)\s*блок", re.IGNORECASE)
OUT_MARK_RE = re.compile(r"\s*❌+\s*")
PER_PIECE_RE = re.compile(r"1\s*шт\s*[-–—:]\s*([\d\s.,]+)", re.IGNORECASE)
PACK_SIZE_RE = re.compile(r"пачк\w*\s*(\d+)\s*шт", re.IGNORECASE)
RESISTANCE_HEADER = ("сопротив", "сопратив")


def _text(v) -> str:
    return "" if v is None else str(v).strip()


def _number(v) -> float | None:
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(_text(v).replace(" ", "").replace(",", "."))
    except ValueError:
        return None


def _tier(label: str, block_size: int | None) -> int | None:
    """Порог уровня в штуках товара (или пачках для расходки)"""
    text = label.strip()
    if m := QTY_TIER_RE.match(text):
        return int(m.group(1))
    if m := BLOCK_TIER_RE.search(text):
        return int(m.group(1))
    if m := BLOCKS_RE.match(text):
        return int(m.group(1)) * (block_size or 10)
    if m := PACK_TIER_RE.match(text):
        return int(m.group(1))
    return None


def _header(ws) -> dict | None:
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=10), start=1):
        texts = [_text(c.value).lower() for c in row]
        name_col = next((i for i, t in enumerate(texts) if t.startswith(NAME_HEADERS)), None)
        if name_col is None:
            continue
        block_size = next((int(m.group(1)) for t in texts if (m := BLOCK_TIER_RE.search(t))), None)
        tiers = {i: _tier(_text(c.value), block_size) for i, c in enumerate(row) if _tier(_text(c.value), block_size)}
        if not tiers:
            continue
        variant_col = next((i for i, t in enumerate(texts) if any(t.startswith(h) for h in VARIANT_HEADERS)), None)
        resistance = variant_col is not None and any(texts[variant_col].startswith(h) for h in RESISTANCE_HEADER)
        return {"row": r, "name": name_col, "variants": variant_col, "tiers": tiers, "resistance": resistance}
    return None


def is_simple_price(content: bytes) -> bool:
    try:
        wb = openpyxl.load_workbook(BytesIO(content), read_only=True)
        return any(_header(ws) for ws in wb.worksheets[:3])
    except Exception:
        return False


def parse_simple(content: bytes, with_images: bool = True) -> ParseResult:
    wb = openpyxl.load_workbook(BytesIO(content))
    sheets: list[ParsedSheet] = []
    warnings: list[str] = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        header = _header(ws)
        if header is None:
            continue
        images: dict[int, bytes] = {}
        if with_images:
            for img in ws._images:
                try:
                    images.setdefault(img.anchor._from.row + 1, img._data())
                except Exception:
                    continue
        title = _text(ws.cell(row=1, column=1).value) or ws.title
        products: list[ParsedProduct] = []
        for r, row in enumerate(ws.iter_rows(min_row=header["row"] + 1), start=header["row"] + 1):
            raw_name = _text(row[header["name"]].value) if header["name"] < len(row) else ""
            if not raw_name:
                continue
            out = "❌" in raw_name
            name = re.sub(r"\s+", " ", OUT_MARK_RE.sub(" ", raw_name)).strip()
            prices: dict[int, int] = {}
            pack_size: int | None = None
            per_piece: float | None = None
            for col, threshold in header["tiers"].items():
                if col >= len(row):
                    continue
                value = row[col].value
                number = _number(value)
                if number is None and isinstance(value, str):
                    # «1 шт - 720 (Пачка 4 шт)»: цена за штуку, продаётся пачками
                    if (m := PER_PIECE_RE.search(value)) and (p := PACK_SIZE_RE.search(value)):
                        per_piece = float(m.group(1).replace(" ", "").replace(",", "."))
                        pack_size = int(p.group(1))
                        number = per_piece * pack_size
                if number and number > 0:
                    prices[threshold] = round(number * 100)
            if pack_size:
                # остальные уровни у расходки заданы за штуку — пересчитываем в цену пачки
                first = min(prices)
                for threshold, amount in list(prices.items()):
                    if threshold != first and amount < prices[first] / pack_size * 1.5:
                        prices[threshold] = amount * pack_size
            if not prices:
                warnings.append(f"{ws.title}, строка {r}: «{name}» — нет цены, пропущен")
                continue
            variants: list[ParsedVariant] = []
            if header["variants"] is not None and header["variants"] < len(row):
                cell = _text(row[header["variants"]].value)
                for line in re.split(r"[\n/]+", cell):
                    v = re.sub(r"\s+", " ", line).strip(" ,;.")
                    if not v:
                        continue
                    if header["resistance"] and re.fullmatch(r"\d+(?:[.,]\d+)?", v):
                        v = f"{v.replace('.', ',')} Ом"
                    if v.lower() not in {x.name.lower() for x in variants}:
                        variants.append(ParsedVariant(name=v, out=out))
            description = None
            if pack_size:
                description = (
                    f"Продаётся пачками по {pack_size} шт — цена указана за пачку "
                    f"({per_piece:g} ₽ за штуку). Количество в заявке — в пачках."
                )
            products.append(ParsedProduct(
                sheet=title, row=r, name=name, brand=None, prices=prices, variants=variants,
                description=description, out=out, image=images.get(r),
            ))
        if products:
            sheets.append(ParsedSheet(title=title, products=products))
    return ParseResult(sheets=sheets, warnings=warnings, tenant_defaults={"price_basis": "qty"})
