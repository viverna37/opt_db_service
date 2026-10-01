"""
Парсер табличных прайсов-выгрузок из учётных систем (МойСклад и похожие,
формат прайса «Галактики»). Чистый модуль — без БД и FastAPI.

Как устроен файл:
- строка заголовков: «Тип», «Код товара модификации», «Наименование»,
  «Ед.изм.», уровни цен «От 20 тыс» / «От 50 тыс» / «От 1 млн» (по сумме
  заявки);
- строка-группа: пустой «Тип», в «Наименовании» путь
  «Товары/1. ЖЕЛЕЗО (Наборы)/1. GEEKVAPE/Aegis Force» -> категории, бренд
  (первое латинское звено пути), модель;
- «Модификация» — вариант товара: товары с одинаковым кодом склеиваются,
  вариант — текст в последних скобках «Набор … Kit (Carbon Black)»;
- «Товар» — самостоятельный товар; позиции одной группы, отличающиеся
  только сопротивлением («GTX 0,15 Ом Mesh Coil» / «GTX 0,2 Ом …»),
  склеиваются в товар с вариантами по сопротивлению;
- нулевые цены (промо, «не для продажи») пропускаются;
- шапка над таблицей: условия (минимальная сумма заказа и т.п.).
"""
from __future__ import annotations

import re
from io import BytesIO

from app.importer.block_price import ParsedProduct, ParsedSheet, ParsedVariant, ParseResult

TYPE_HEADERS = {"тип"}
NAME_HEADERS = {"наименование"}
CODE_HEADERS = {"код товара модификации", "код", "артикул", "код товара"}
TIER_HEADER_RE = re.compile(r"^от\s*([\d\s.,]+)\s*(тыс|млн|₽|руб|р)?\.?$", re.IGNORECASE)
# «ОПТ 5к» / «ОПТ 25к» / «ОПТ 150к» — уровни по сумме заявки в тысячах
OPT_K_RE = re.compile(r"^опт\s*([\d.,]+)\s*(к|тыс|млн)\.?$", re.IGNORECASE)
# «Мелкий ОПТ» / «Крупный ОПТ» — уровни без порогов в прайсе: ставим условные пороги по сумме
# заявки (крупный — от 40 000 ₽, как «Поды и расходники от 40т»), владелец поправит в админке
NAMED_TIERS = {"мелкий опт": (1, "Мелкий опт"), "крупный опт": (40000, "Крупный опт от 40 000 ₽")}
STOCK_HEADERS = {"остаток", "остатки", "в наличии", "наличие"}
# Пометки в конце названия — не вкус: «… (Арбуз) (топ)», «… (акция)»
TAG_RE = re.compile(r"\s*\((топ|акция|хит|sale|new|новинка[^()]*)\)\s*$", re.IGNORECASE)
TAG_LABELS = {"топ": "Топ продаж", "акция": "Акция", "хит": "Хит", "sale": "Акция", "new": "Новинка"}
EMOJI_RE = re.compile("[\U0001F000-\U0001FFFF\u2600-\u27BF\uFE0F]")
# «a/b» — разделитель пути; «Испарители / Картриджи» (с пробелами) — часть названия группы
PATH_SPLIT_RE = re.compile(r"(?<!\s)/(?!\s)")
HOT_RE = re.compile(r"\s*[-–—]\s*горячее предложение!?\s*$", re.IGNORECASE)
NUMBER_PREFIX_RE = re.compile(r"^\s*\d+\.\s*")
NEW_MARK_RE = re.compile(r"\s*\((?:новинка|new)[^()]*\)\s*", re.IGNORECASE)
LATIN_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9 &.\-+']*$")
RESISTANCE_RE = re.compile(r"(\d+(?:[.,]\d+)?)\s*(Ω|Ом|ом|ohm|om)(?![A-Za-zА-Яа-я])", re.IGNORECASE)
SKIP_ROOTS = {"промо (мерч)", "промо"}
MIN_ORDER_RE = re.compile(r"минимальн\w*\s+сумм\w*\s+заказа\s*[—\-:]*\s*([\d\s]+)", re.IGNORECASE)
TG_RE = re.compile(r"https?://t\.me/([A-Za-z0-9_]{4,})")
MOYSKLAD_RE = re.compile(r"https?://b2b\.moysklad\.ru/public/([A-Za-z0-9_\-]+)")


# ---------- чтение файла в «сетку» значений ----------

def _grid_xls(content: bytes) -> tuple[list[list], list[str]]:
    import xlrd

    book = xlrd.open_workbook(file_contents=content)
    sheet = book.sheet_by_index(0)
    grid = [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
    links = [link.url_or_path for link in getattr(sheet, "hyperlink_list", [])]
    return grid, links


def _grid_xlsx(content: bytes) -> tuple[list[list], list[str]]:
    import openpyxl

    wb = openpyxl.load_workbook(BytesIO(content), read_only=False)
    ws = wb.worksheets[0]
    grid = [[c.value if c.value is not None else "" for c in row] for row in ws.iter_rows()]
    links = [c.hyperlink.target for row in ws.iter_rows() for c in row if c.hyperlink and c.hyperlink.target]
    return grid, links


def read_grid(content: bytes, filename: str) -> tuple[list[list], list[str]]:
    return _grid_xls(content) if filename.lower().endswith(".xls") else _grid_xlsx(content)


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _number(value) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(_text(value).replace(" ", "").replace(",", "."))
    except ValueError:
        return None


def _threshold(label: str) -> int | None:
    """«От 20 тыс» -> 20000, «От 1 млн» -> 1000000, «От 3 000 ₽» -> 3000, «ОПТ 5к» -> 5000"""
    if named := NAMED_TIERS.get(re.sub(r"\s+", " ", label.strip().lower())):
        return named[0]
    m = TIER_HEADER_RE.match(label.strip()) or OPT_K_RE.match(label.strip())
    if not m:
        return None
    number = float(m.group(1).replace(" ", "").replace(",", "."))
    unit = (m.group(2) or "").lower()
    return int(number * (1000 if unit in ("тыс", "к") else 1_000_000 if unit == "млн" else 1))


def find_header(grid: list[list]) -> dict | None:
    """Строка заголовков таблицы и номера колонок; None — это не табличный прайс"""
    for r, row in enumerate(grid[:40]):
        texts = [_text(v).lower().strip(" .:") for v in row]
        name_cols = [c for c, t in enumerate(texts) if t in NAME_HEADERS]
        tiers = {c: _threshold(_text(v)) for c, v in enumerate(row) if _threshold(_text(v))}
        if name_cols and tiers:
            # колонка «Тип» и код — в строке заголовков экспорта (обычно первая строка файла)
            type_col = code_col = None
            for rr in range(0, r + 1):
                for c, t in enumerate(_text(v).lower() for v in grid[rr]):
                    if t in TYPE_HEADERS and type_col is None:
                        type_col = c
                    if t in CODE_HEADERS and code_col is None:
                        code_col = c
            stock_col = next((c for c, t in enumerate(texts) if t in STOCK_HEADERS), None)
            return {"row": r, "name": name_cols[-1], "tiers": tiers, "type": type_col, "code": code_col, "stock": stock_col}
    return None


def is_table_price(content: bytes, filename: str) -> bool:
    try:
        grid, _ = read_grid(content, filename)
    except Exception:
        return False
    return find_header(grid) is not None


# ---------- разбор ----------

def _clean_segment(text: str) -> tuple[str, bool]:
    hot = bool(HOT_RE.search(text))
    text = EMOJI_RE.sub("", text)
    text = HOT_RE.sub("", NUMBER_PREFIX_RE.sub("", text)).strip()
    text = re.sub(r"\s+от\s+\d+\s*т\.?$", "", text)  # «Поды и расходники от 40т»
    return re.sub(r"\s+", " ", text), hot


def _strip_tags(name: str) -> tuple[str, list[str]]:
    """«… (Арбуз) (топ)» -> («… (Арбуз)», ["Топ продаж"])"""
    tags = []
    while m := TAG_RE.search(name):
        tags.append(TAG_LABELS.get(m.group(1).lower().split()[0], m.group(1)))
        name = name[: m.start()].rstrip()
    return name, tags


def _category_name(segment: str) -> str:
    """«ЖЕЛЕЗО (Наборы)» -> «Железо (наборы)», «СМЕННЫЙ КАРТРИДЖ» -> «Сменный картридж»"""
    letters = [ch for ch in segment if ch.isalpha()]
    upper = sum(1 for ch in letters if ch.isupper())
    if letters and upper / len(letters) >= 0.5:
        # «SALT РФ» -> «Salt РФ»: короткие аббревиатуры (до 3 букв) оставляем как есть
        words = [w if len(w) <= 3 and w.isupper() and i else w.lower() for i, w in enumerate(segment.split(" "))]
        text = " ".join(words)
        return text[:1].upper() + text[1:]
    return segment


def _split_last_parens(name: str) -> tuple[str, str | None]:
    """«Набор X Kit (Lightning Yellow (Новинка 06.26))» -> («Набор X Kit», «Lightning Yellow (Новинка 06.26)»)"""
    text = name.rstrip()
    if not text.endswith(")"):
        return text, None
    depth = 0
    for i in range(len(text) - 1, -1, -1):
        if text[i] == ")":
            depth += 1
        elif text[i] == "(":
            depth -= 1
            if depth == 0:
                return text[:i].rstrip(), text[i + 1:-1].strip()
    return text, None


def _clean_name(text: str) -> tuple[str, bool]:
    """Убирает «(новинка 07.26)» из названия; второй элемент — была ли пометка"""
    is_new = bool(NEW_MARK_RE.search(text))
    return re.sub(r"\s+", " ", NEW_MARK_RE.sub(" ", text)).strip(), is_new


def parse_table(content: bytes, filename: str) -> ParseResult:
    grid, links = read_grid(content, filename)
    header = find_header(grid)
    if header is None:
        raise ValueError("Не нашёл строку заголовков с «Наименование» и уровнями цен «От …»")
    warnings: list[str] = []
    name_col, type_col, code_col, tier_cols = header["name"], header["type"], header["code"], header["tiers"]
    stock_col = header.get("stock")
    named_tiers = {
        NAMED_TIERS[key][0]: NAMED_TIERS[key][1]
        for c in tier_cols if (key := re.sub(r"\s+", " ", _text(grid[header["row"]][c]).lower())) in NAMED_TIERS
    }

    # Шапка над таблицей: условия и контакты
    head_text = "\n".join(_text(v) for row in grid[: header["row"]] for v in row if _text(v))
    defaults: dict = {}
    if m := MIN_ORDER_RE.search(head_text):
        defaults["min_order_amount"] = int(re.sub(r"\s", "", m.group(1)))
    elif tier_cols and not named_tiers:
        # «ОПТ 5к / 25к / 150к»: ниже младшего уровня опта нет — он и есть минимальная заявка
        defaults["min_order_amount"] = min(tier_cols.values())
    tg = [m.group(1) for link in links if (m := TG_RE.search(link or ""))]
    if tg:
        defaults["manager_username"] = tg[0]
    moysklad = [m.group(1) for link in links if (m := MOYSKLAD_RE.search(link or ""))]
    if moysklad:
        defaults["moysklad_public_id"] = moysklad[0]  # публичный B2B-каталог: фото и остатки
    conditions = [line.strip(" -•\n") for line in head_text.splitlines() if line.strip().startswith("-")]
    if conditions:
        defaults["welcome_text"] = "\n".join(f"• {c}" for c in conditions[:6])

    groups: list[dict] = []  # {path, brand, hot, rows}
    current: dict | None = None
    def cell(row, col) -> str:
        return _text(row[col]) if col is not None and col < len(row) else ""

    for r in range(header["row"] + 1, len(grid)):
        row = grid[r]
        name = cell(row, name_col)
        row_type = cell(row, type_col)
        prices = {}
        for col, threshold in tier_cols.items():
            value = _number(row[col]) if col < len(row) else None
            if value and value > 0:
                prices[threshold] = round(value * 100)
        # строка-группа: путь «Товары/1. ЖЕЛЕЗО/…» — в «Наименовании» или (1С-выгрузки) в колонке кода
        path = name if (not row_type and "/" in name and not prices) else ""
        # 1С без колонки «Тип» и кода: группа верхнего уровня без «/» («Поды», «Мистери Бокс»)
        if (not path and name and not prices and type_col is None and code_col is None
                and "(" not in name and len(name) <= 60):
            path = name
        if not path and not name and "/" in cell(row, code_col):
            path = cell(row, code_col)
        if path:
            segments = [s for s in PATH_SPLIT_RE.split(path) if s.strip()]
            if segments and segments[0].strip().lower() in {"товары", "products"}:
                segments = segments[1:]
            cleaned = [_clean_segment(s) for s in segments]
            current = {"segments": [c[0] for c in cleaned if c[0]], "hot": any(c[1] for c in cleaned), "rows": []}
            groups.append(current)
            continue
        if not name:
            continue
        if current is None:
            current = {"segments": [], "hot": False, "rows": []}
            groups.append(current)
        code = cell(row, code_col)
        stock = _number(row[stock_col]) if stock_col is not None and stock_col < len(row) else None
        name, tags = _strip_tags(EMOJI_RE.sub("", name).strip())
        if type_col is None:
            # без колонки «Тип»: вкус/цвет — в последних скобках названия
            row_type = "модификация" if name.endswith(")") else "товар"
        current["rows"].append({
            "r": r + 1, "type": row_type.lower(), "name": name, "code": code or None, "prices": prices,
            "stock": int(stock) if stock is not None else None, "tags": tags, "code_is_variant": type_col is None,
        })

    brand_names: dict[str, str] = {}
    # Раздел, где подгруппы почти все латинские бренды («Жидкости/Angry Ape», «Жидкости/BJORN»),
    # — кириллические подгруппы там тоже бренды («Жидкости/Злая Монашка»), а не подкатегории
    children: dict[str, set[str]] = {}
    for group in groups:
        if len(group["segments"]) >= 2:
            children.setdefault(group["segments"][0].lower(), set()).add(group["segments"][1])
    brand_parents = {
        parent for parent, names in children.items()
        if len(names) >= 5 and sum(bool(LATIN_RE.match(n)) for n in names) / len(names) >= 0.7
    }
    sheets: dict[str, ParsedSheet] = {}
    for group in groups:
        segments = group["segments"]
        if segments and segments[0].lower() in SKIP_ROOTS:
            continue
        brand_idx = next((i for i, s in enumerate(segments) if i > 0 and LATIN_RE.match(s)), None)
        if brand_idx is None and len(segments) == 2 and segments[0].lower() in brand_parents:
            brand_idx = 1
        brand = None
        if brand_idx is not None:
            key = segments[brand_idx].lower().replace(" ", "")
            brand = brand_names.setdefault(key, segments[brand_idx])
        category_path = [_category_name(s) for s in (segments[:brand_idx] if brand_idx is not None else segments)][:3]
        if not category_path:
            category_path = ["Каталог"]
        sheet = sheets.setdefault(category_path[0], ParsedSheet(title=category_path[0], products=[]))
        for product in _group_products(group, brand, category_path, warnings):
            sheet.products.append(product)

    # одинаковые названия в одной категории — разводим номером
    for sheet in sheets.values():
        seen: dict[tuple, int] = {}
        for p in sheet.products:
            key = (tuple(p.category_path or []), p.name.lower())
            seen[key] = seen.get(key, 0) + 1
            if seen[key] > 1:
                p.name = f"{p.name} #{seen[key]}"
    if named_tiers:
        defaults["tier_labels"] = named_tiers
    for sheet in sheets.values():
        for product in sheet.products:
            for variant in product.variants:
                variant.name = re.sub(r"^\((.*)\)$", r"\1", variant.name).strip()
    return ParseResult(sheets=[s for s in sheets.values() if s.products], warnings=warnings, tenant_defaults=defaults)


def _group_products(group: dict, brand: str | None, category_path: list[str], warnings: list[str]) -> list[ParsedProduct]:
    products: list[ParsedProduct] = []
    rows = [row for row in group["rows"] if row["prices"]]
    for row in group["rows"]:
        if not row["prices"]:
            warnings.append(f"строка {row['r']}: «{row['name']}» — нулевая цена, пропущена")

    def make(name: str, sku: str | None, row: dict) -> ParsedProduct:
        clean, is_new = _clean_name(name)
        notes = []
        if is_new:
            notes.append("Новинка")
        if group["hot"]:
            notes.append("Горячее предложение")
        notes.extend(t for t in row.get("tags", []) if t not in notes)
        return ParsedProduct(
            sheet=category_path[0], row=row["r"], name=clean, brand=brand, prices=dict(row["prices"]),
            description=" · ".join(notes) or None, sku=sku, category_path=category_path, source_name=row["name"],
            stock_qty=row.get("stock"),
        )

    # «Модификация»: склейка по коду (или по названию без скобок)
    by_key: dict[str, ParsedProduct] = {}
    for row in [r for r in rows if r["type"].startswith("модиф")]:
        base, variant = _split_last_parens(row["name"])
        # код — общий у модификаций товара (МойСклад) или свой у каждой строки (1С): тогда склеиваем по названию
        key = base.lower() if row.get("code_is_variant") or not row["code"] else row["code"]
        product = by_key.get(key)
        if product is None:
            product = make(base, None if row.get("code_is_variant") else row["code"], row)
            by_key[key] = product
            products.append(product)
        else:
            for tag in row.get("tags", []):
                if tag not in (product.description or ""):
                    product.description = " · ".join(filter(None, [product.description, tag]))
        if variant:
            vname, _ = _clean_name(variant)
            vname = re.sub(r"^new colou?r\s+", "", vname, flags=re.IGNORECASE)
            product.variants.append(ParsedVariant(
                name=vname or variant, prices=dict(row["prices"]), source_name=row["name"], stock_qty=row.get("stock"),
                sku=row["code"] if row.get("code_is_variant") else None,
            ))

    # «Товар»: склейка позиций, отличающихся только сопротивлением
    singles = [r for r in rows if not r["type"].startswith("модиф")]
    buckets: dict[str, list[tuple[dict, str]]] = {}
    order: list[str] = []
    for row in singles:
        m = RESISTANCE_RE.search(row["name"])
        if m:
            base = re.sub(r"\s+", " ", (row["name"][: m.start()] + " " + row["name"][m.end():])).strip()
            key = base.lower()
            token = m.group(0).strip()
        else:
            key, token = f"__single_{row['r']}", ""
        if key not in buckets:
            order.append(key)
        buckets.setdefault(key, []).append((row, token))
    for key in order:
        items = buckets[key]
        if len(items) == 1:
            products.append(make(items[0][0]["name"], items[0][0]["code"], items[0][0]))
            continue
        first_row = items[0][0]
        base = re.sub(r"\s+", " ", (first_row["name"][: RESISTANCE_RE.search(first_row["name"]).start()]
                                   + " " + first_row["name"][RESISTANCE_RE.search(first_row["name"]).end():])).strip()
        base = re.sub(r"\s+([,)])", r"\1", base)
        product = make(base, None, first_row)
        for row, token in items:
            product.variants.append(ParsedVariant(
                name=token, prices=dict(row["prices"]), source_name=row["name"], stock_qty=row.get("stock"),
            ))
        products.append(product)

    # цена товара — цена первого варианта; варианты с такой же ценой не храним отдельно
    for product in products:
        if product.variants:
            product.prices = dict(product.variants[0].prices or product.prices)
            for v in product.variants:
                if v.prices == product.prices:
                    v.prices = None
    return products
