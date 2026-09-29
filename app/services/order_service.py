"""Заявки: отправка из корзины (снимок + пересчёт цен), повтор, смена статуса, текст и Excel."""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font

from app.database.models import (
    Order,
    OrderItem,
    OrderStatus,
    Tenant,
    TenantUser,
)
from app.database.repository.main_repository import Repository
from app.services.cart_service import build_cart_view
from app.services.catalog_service import is_orderable
from app.utils.dt import as_aware_utc, utcnow
from app.utils.money import format_money, safe_zone

STATUS_LABELS = {
    OrderStatus.new: "Новая",
    OrderStatus.in_progress: "В работе",
    OrderStatus.done: "Выполнена",
    OrderStatus.cancelled: "Отменена",
}


class OrderError(Exception):
    def __init__(self, code: str, message: str, data: dict | None = None):
        super().__init__(message)
        self.code = code
        self.message = message
        self.data = data or {}


BLOCKER_MESSAGES = {
    "empty": "Корзина пуста",
    "unavailable_items": "В корзине есть позиции, которых нет в наличии — уберите их",
    "no_price_items": "В корзине есть позиции без цены — уточните у менеджера",
    "below_min_amount": "Сумма заявки меньше минимальной",
}


async def submit_order(repo: Repository, tenant: Tenant, member: TenantUser, comment: str | None) -> Order:
    """
    Цена всегда пересчитывается здесь, в момент отправки — данным с фронта
    не доверяем (фронт вообще не присылает цены). Позиции копируются
    снимком: товар потом могут переименовать или удалить.
    """
    cart = await repo.cart.get_for_member(member.id)
    view = await build_cart_view(repo, tenant, cart)
    if view.response.blockers:
        first = view.response.blockers[0]
        raise OrderError(first, BLOCKER_MESSAGES.get(first, "Нельзя отправить заявку"),
                         {"blockers": view.response.blockers})

    number = await repo.tenant.next_order_number(tenant.id)
    now = utcnow()
    order = Order(
        tenant_id=tenant.id, number=number, tenant_user_id=member.id, status=OrderStatus.new,
        comment=(comment or "").strip() or None, total=view.response.total, created_at=now, updated_at=now,
    )
    items = [
        OrderItem(
            product_id=line.product.id,
            variant_id=line.variant.id,
            product_name=line.product.name,
            variant_name=None if line.variant.is_default else line.variant.name,
            sku=line.variant.sku or line.product.sku,
            qty=line.qty,
            price=line.unit_price,
            tier_label=line.tier_label,
            amount=line.amount,
        )
        for line in view.priced_lines
    ]
    order = await repo.order.create(order, items, actor_id=member.id)
    await repo.cart.clear(cart)
    return order


async def repeat_order(repo: Repository, tenant: Tenant, member: TenantUser, order: Order) -> tuple[int, list[str]]:
    """«Повторить»: кладёт позиции снова в корзину (прибавляя к тому, что там уже есть)"""
    cart = await repo.cart.get_or_create(tenant.id, member.id)
    in_cart = {item.variant_id: item.qty for item in cart.items}
    variants = await repo.variant.get_many(tenant.id, [i.variant_id for i in order.items if i.variant_id])
    products = await repo.product.get_many(tenant.id, list({v.product_id for v in variants.values()}))
    added, skipped = 0, []
    for item in order.items:
        variant = variants.get(item.variant_id) if item.variant_id else None
        product = products.get(variant.product_id) if variant else None
        if variant is None or not is_orderable(variant, product):
            skipped.append(f"{item.product_name} — {item.variant_name}" if item.variant_name else item.product_name)
            continue
        await repo.cart.set_qty(cart, variant.id, in_cart.get(variant.id, 0) + item.qty)
        in_cart[variant.id] = in_cart.get(variant.id, 0) + item.qty
        added += 1
    return added, skipped


async def change_status(repo: Repository, order: Order, new_status: OrderStatus, actor: TenantUser) -> bool:
    """False — статус и так такой (повторное нажатие «Взять в работу» в боте не плодит историю)"""
    if order.status == new_status:
        return False
    await repo.order.add_history(order.id, order.status, new_status, actor.id)
    order.status = new_status
    order.updated_at = utcnow()
    return True


# ---------- Текст и Excel ----------

def _client_label(member: TenantUser) -> str:
    user = member.tg_user
    name = " ".join(filter(None, [user.first_name, user.last_name])) or f"id {user.telegram_id}"
    label = f"{name} (@{user.username})" if user.username else name
    return f"{label}, {member.note}" if member.note else label


def _local(tenant: Tenant, value: datetime) -> datetime:
    return as_aware_utc(value).astimezone(safe_zone(tenant.timezone))


def order_text(tenant: Tenant, order: Order, client: TenantUser, include_client: bool = True) -> str:
    """«Скопировать текстом» (для 1С/накладной) — он же текст уведомлений в бот"""
    created = _local(tenant, order.created_at).strftime("%d.%m.%Y %H:%M")
    lines = [f"Заявка №{order.number} от {created}"]
    if include_client:
        lines.append(f"Клиент: {_client_label(client)}")
    lines.append("")
    for index, item in enumerate(order.items, start=1):
        name = f"{item.product_name} — {item.variant_name}" if item.variant_name else item.product_name
        sku = f" [{item.sku}]" if item.sku else ""
        lines.append(
            f"{index}. {name}{sku}: {item.qty} шт × {format_money(item.price, tenant.currency)}"
            f" = {format_money(item.amount, tenant.currency)}"
        )
    lines.append("")
    lines.append(f"Итого: {format_money(order.total, tenant.currency)}")
    if order.comment:
        lines.append(f"Комментарий: {order.comment}")
    return "\n".join(lines)


def orders_xlsx(tenant: Tenant, orders: list[Order]) -> bytes:
    """Экспорт одной или нескольких заявок: строка на позицию, суммы в рублях числами"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Заявки"
    header = ["№ заявки", "Дата", "Статус", "Клиент", "Товар", "Вариант", "Артикул", "Кол-во", "Цена", "Уровень",
              "Сумма", "Комментарий"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for order in orders:
        client = _client_label(order.tenant_user)
        created = _local(tenant, order.created_at).replace(tzinfo=None)
        for item in order.items:
            ws.append([
                order.number, created, STATUS_LABELS[order.status], client, item.product_name, item.variant_name,
                item.sku, item.qty, item.price / 100, item.tier_label, item.amount / 100, order.comment,
            ])
    for column, width in zip("ABCDEFGHIJKL", (10, 17, 12, 30, 35, 22, 14, 8, 10, 12, 12, 30)):
        ws.column_dimensions[column].width = width
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        for cell in row:
            cell.number_format = "DD.MM.YYYY HH:MM"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
